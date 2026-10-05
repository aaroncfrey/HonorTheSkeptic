#!/usr/bin/env python3
"""
Build the Honor the Skeptic site (and handouts) from the files in content/.

  python3 tools/build.py              # live site -> public/   (only lessons with status: published)
  python3 tools/build.py --preview    # preview  -> preview/   (draft + published; never deployed)
  python3 tools/build.py --no-handouts   # skip PDF/Word generation (faster)

Lesson status values (in content/lessons/week-N.yml):
  planned   -> shown on the home page as "Coming <date>", no page
  draft     -> visible only in --preview builds
  published -> on the live site, with handouts

Requires: PyYAML, Jinja2, Markdown, segno, python-docx; Playwright + Chromium for PDFs.
"""
import argparse, datetime as dt, html, io, re, shutil, sys, urllib.parse
from pathlib import Path

import markdown as mdlib
import segno
import yaml
from jinja2 import Environment, FileSystemLoader, select_autoescape
from markupsafe import Markup

ROOT = Path(__file__).resolve().parent.parent
CONTENT = ROOT / "content"
TEMPLATES = ROOT / "tools" / "templates"
PUBLIC = ROOT / "public"

# ---------------------------------------------------------------- helpers
def md(text):
    if not text:
        return Markup("")
    return Markup(mdlib.markdown(str(text).strip(), extensions=["smarty", "sane_lists"]))

def mdi(text):
    """Inline markdown (no wrapping <p>)."""
    h = str(md(text)).strip()
    h = re.sub(r"^<p>(.*)</p>$", r"\1", h, flags=re.S)
    return Markup(h)

def secs(t):
    if t in (None, ""):
        return None
    if isinstance(t, int):
        return t
    parts = [int(p) for p in str(t).split(":")]
    s = 0
    for p in parts:
        s = s * 60 + p
    return s

def fmt_len(a, b):
    if a is None or b is None or b <= a:
        return None
    d = b - a
    m, s = divmod(d, 60)
    return f"{m}:{s:02d}" if m else f"{s} sec"

def slugify(s):
    s = re.sub(r"[^\w\s-]", "", str(s).lower())
    return re.sub(r"[\s_]+", "-", s).strip("-")[:48] or "section"

def bible_link(ref, version):
    q = urllib.parse.quote_plus(ref.replace("–", "-").replace("—", "-"))
    return f"https://www.biblegateway.com/passage/?search={q}&version={version}"

def date_label(d):
    if not d:
        return None
    if isinstance(d, str):
        try:
            d = dt.date.fromisoformat(d)
        except ValueError:
            return d
    return f"{d.strftime('%A')}, {d.strftime('%B')} {d.day}"

def qr_svg(url, color="#1c2230"):
    return segno.make(url, error="m").svg_inline(scale=4, border=0, dark=color, omitsize=True)

def qr_png(url):
    buf = io.BytesIO()
    segno.make(url, error="m").save(buf, kind="png", scale=8, border=1)
    buf.seek(0)
    return buf

def yt_url(vid, start):
    return f"https://youtu.be/{vid}" + (f"?t={start}" if start else "")

def label_for(s):
    if s.get("heading"):
        return s["heading"]
    return {"clip": "Video: " + s.get("title", ""), "question": "Question", "questions": "Discussion",
            "scripture": "Scripture", "quote": "Quote", "activity": "Activity"}.get(s["type"], s["type"].title())

# ---------------------------------------------------------------- load
def load(preview):
    site = yaml.safe_load((CONTENT / "site.yml").read_text(encoding="utf-8"))
    lessons = []
    for f in sorted(CONTENT.glob("lessons/week-*.yml"), key=lambda p: int(re.findall(r"\d+", p.stem)[0])):
        L = yaml.safe_load(f.read_text(encoding="utf-8")) or {}
        L.setdefault("week", int(re.findall(r"\d+", f.stem)[0]))
        L["slug"] = f"week-{L['week']}"
        L["status"] = str(L.get("status", "planned")).lower()
        L["live"] = L["status"] == "published" or (preview and L["status"] == "draft")
        L["date_label"] = date_label(L.get("date"))
        L["sections"] = L.get("sections") or []
        L["downloads"] = L.get("downloads") or []
        used = set()
        contents = []
        for i, s in enumerate(L["sections"]):
            s.setdefault("type", "text")
            base = slugify(s.get("id") or s.get("heading") or s.get("title") or f"{s['type']}-{i+1}")
            sid, n = base, 2
            while sid in used:
                sid, n = f"{base}-{n}", n + 1
            used.add(sid)
            s["id"] = sid
            if s["type"] == "scripture":
                for p in s.get("passages") or []:
                    p["link"] = bible_link(p["ref"], site.get("bible_version", "ESV"))
            if s["type"] == "clip":
                vid = s.get("youtube")
                if not vid:
                    sys.exit(f"{f.name}: clip '{s.get('title')}' needs a youtube: video id")
                s["start_s"] = secs(s.get("start")) or 0
                s["end_s"] = secs(s.get("end"))
                s["start"] = s.get("start") or "0:00"
                s["length"] = fmt_len(s["start_s"], s["end_s"])
                s["url"] = yt_url(vid, s["start_s"])
                s["qr_svg"] = Markup(qr_svg(s["url"]))
                s["ties_to_links"] = [{"ref": r, "link": bible_link(r, site.get("bible_version", "ESV"))} for r in s.get("ties_to") or []]
            contents.append({"id": sid, "label": label_for(s)})
        L["contents"] = contents
        base = f"handouts/doubters-welcome-week-{L['week']}"
        L["handout_base"] = base
        L["handouts"] = [{"label": "Handout (PDF)", "path": base + ".pdf"},
                         {"label": "Handout (Word)", "path": base + ".docx"}]
        lessons.append(L)
    live = [l for l in lessons if l["live"]]
    for i, l in enumerate(live):
        l["prev"] = live[i - 1] if i > 0 else None
        l["next"] = live[i + 1] if i + 1 < len(live) else None
    return site, lessons

# ---------------------------------------------------------------- render site
def env():
    e = Environment(loader=FileSystemLoader(str(TEMPLATES)), autoescape=select_autoescape(["html"]),
                    trim_blocks=True, lstrip_blocks=True)
    e.filters["md"] = md
    e.filters["mdi"] = mdi
    return e

def build_site(site, lessons, out, preview):
    E = env()
    build_id = dt.datetime.now().strftime("%Y%m%d%H%M")
    common = dict(site=site, lessons_all=lessons, nav_lessons=lessons, preview=preview, build_id=build_id)

    # clear generated lesson folders (keeps assets/, files/, handouts/)
    for d in out.glob("week-*"):
        shutil.rmtree(d)
    if preview:
        if (out / "assets").exists():
            shutil.rmtree(out / "assets")
        shutil.copytree(PUBLIC / "assets", out / "assets")
        if (PUBLIC / "files").exists():
            if (out / "files").exists():
                shutil.rmtree(out / "files")
            shutil.copytree(PUBLIC / "files", out / "files")

    (out / "index.html").write_text(E.get_template("index.html").render(root="", lesson=None, **common), encoding="utf-8")
    for L in lessons:
        if not L["live"]:
            continue
        d = out / L["slug"]
        d.mkdir(parents=True, exist_ok=True)
        (d / "index.html").write_text(E.get_template("lesson.html").render(root="../", lesson=L, **common), encoding="utf-8")
        print(f"  page   {d.relative_to(ROOT)}/index.html  [{L['status']}]")

    # remove handouts for lessons that are no longer live
    hdir = out / "handouts"
    if hdir.exists():
        live_bases = {Path(L["handout_base"]).name for L in lessons if L["live"]}
        for f in hdir.iterdir():
            if f.stem not in live_bases:
                f.unlink()

# ---------------------------------------------------------------- handouts
def build_handouts(site, lessons, out):
    E = env()
    hdir = out / "handouts"
    hdir.mkdir(parents=True, exist_ok=True)
    targets = [L for L in lessons if L["live"]]
    if not targets:
        return
    site_qr = Markup(qr_svg(site["url"]))
    asset_base = (PUBLIC / "assets").as_uri() + "/"
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("  ! Playwright not installed: skipping PDFs")
        sync_playwright = None

    if sync_playwright:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page()
            for L in targets:
                h = E.get_template("handout.html").render(site=site, lesson=L, lessons_all=lessons,
                                                          site_qr_svg=site_qr, asset_base=asset_base)
                tmp = ROOT / "tools" / f".handout-{L['slug']}.html"
                tmp.write_text(h, encoding="utf-8")
                page.goto(tmp.as_uri())
                page.wait_for_load_state("networkidle")
                pdf = out / (L["handout_base"] + ".pdf")
                page.pdf(path=str(pdf), format="Letter", print_background=True, prefer_css_page_size=True,
                         display_header_footer=True, header_template="<span></span>",
                         footer_template=f'<div style="width:100%;font:7pt Helvetica,Arial,sans-serif;color:#888;text-align:center">{html.escape(site["class_name"])} · Week {L["week"]} · page <span class="pageNumber"></span> of <span class="totalPages"></span></div>')
                tmp.unlink()
                print(f"  pdf    {pdf.relative_to(ROOT)}")
            browser.close()

    for L in targets:
        path = out / (L["handout_base"] + ".docx")
        build_docx(site, L, lessons, path)
        print(f"  word   {path.relative_to(ROOT)}")

# --- Word (.docx)
def build_docx(site, L, lessons, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    INK, ACC, MUTED, LINE = RGBColor(0x1C, 0x22, 0x30), RGBColor(0x93, 0x37, 0x1F), RGBColor(0x73, 0x78, 0x89), "D9D2C4"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.top_margin = sec.bottom_margin = Inches(0.55)

    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(10)
    st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.08

    def border(p, side="bottom", color=LINE, sz=6):
        pPr = p._p.get_or_add_pPr()
        bdr = pPr.find(qn("w:pBdr"))
        if bdr is None:
            bdr = OxmlElement("w:pBdr")
            pPr.append(bdr)
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "4"); el.set(qn("w:color"), color)
        bdr.append(el)

    def shade(p, fill="F5EBE5"):
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
        pPr.append(shd)

    def widths(tbl, inches):
        tbl.autofit = False
        tblPr = tbl._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        grid = tbl._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), inches):
            gc.set(qn("w:w"), str(int(w * 1440)))
        for row in tbl.rows:
            for c, w in zip(row.cells, inches):
                c.width = Inches(w)

    def runs(p, text, size=None, color=None, italic=False, bold=False, font=None):
        """Add text with **bold** / *italic* markdown support."""
        text = str(text).replace("--", "–")
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))", text):
            if not part:
                continue
            b, i = bold, italic
            if part.startswith("**"):
                part, b = part[2:-2], True
            elif part.startswith("*"):
                part, i = part[1:-1], not italic
            elif part.startswith("["):
                part = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", part)
            r = p.add_run(part)
            r.bold, r.italic = b, i
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
            if font: r.font.name = font
        return p

    def label(text, color=ACC, before=10):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(7.5); r.font.color.rgb = color; r.font.name = "Arial"
        return p

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text); r.bold = True; r.font.size = Pt(13)
        return p

    def paras(text, **kw):
        for chunk in re.split(r"\n\s*\n", str(text).strip()):
            if chunk.strip():
                runs(doc.add_paragraph(), " ".join(chunk.split()), **kw)

    def lines(n):
        for _ in range(n):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.add_run(" ").font.size = Pt(16)
            border(p)

    def numbered(items, note_lines=0):
        for n, q in enumerate(items, 1):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            runs(p, f"{n}.\t", bold=True, color=ACC)
            runs(p, q)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))
            if note_lines:
                lines(note_lines)

    # masthead
    p = doc.add_paragraph()
    runs(p, f"{site['class_name'].upper()} · WEEK {L['week']} OF {len(lessons)}", size=7.5, bold=True, color=ACC, font="Arial")
    runs(p, f"\t{site['church']['name']} · {L['date_label'] or site['church']['town']}", size=7.5, color=MUTED, font="Arial")
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), alignment=2)
    border(p, color="1C2230", sz=12)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    runs(p, L["title"], size=26, bold=False)
    if L.get("subtitle"):
        runs(doc.add_paragraph(), L["subtitle"], size=12, italic=True, color=RGBColor(0x47, 0x4E, 0x5E))
    if L.get("big_idea"):
        p = doc.add_paragraph()
        shade(p)
        border(p, "left", "93371F", 18)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(10)
        runs(p, "BIG IDEA   ", size=7.5, bold=True, color=ACC, font="Arial")
        runs(p, L["big_idea"], size=11.5)

    for s in L["sections"]:
        if s.get("handout") is False:
            continue
        t = s["type"]
        if s.get("kicker"):
            label(s["kicker"])
        if s.get("heading"):
            heading(s["heading"])
        if t == "text":
            paras(s.get("body", ""))
        elif t == "question":
            runs(doc.add_paragraph(), s["text"], size=13, italic=True)
            if s.get("body"): paras(s["body"])
            if s.get("note_lines"): lines(s["note_lines"])
        elif t == "scripture":
            if s.get("body"): paras(s["body"])
            for ps in s.get("passages") or []:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(1.4)
                p.paragraph_format.first_line_indent = Inches(-1.4)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(1.4))
                p.paragraph_format.space_after = Pt(2)
                runs(p, ps["ref"], bold=True, size=9.5, font="Arial")
                if ps.get("note"):
                    runs(p, "\t" + ps["note"], color=RGBColor(0x47, 0x4E, 0x5E))
                if ps.get("text"):
                    q = doc.add_paragraph()
                    q.paragraph_format.left_indent = Inches(0.25)
                    border(q, "left", "A8812F", 8)
                    runs(q, " ".join(str(ps["text"]).split()), size=10)
        elif t == "clip":
            tbl = doc.add_table(rows=1, cols=2)
            widths(tbl, [5.5, 1.4])
            c0, c1 = tbl.rows[0].cells
            p0 = c0.paragraphs[0]
            runs(p0, s["title"], bold=True, size=11)
            meta = " · ".join(x for x in [s.get("speaker"), s.get("channel") if s.get("channel") != s.get("speaker") else None,
                                           s["start"] + (("–" + str(s["end"])) if s.get("end") else "")] if x)
            runs(c0.add_paragraph(), meta, size=8, color=MUTED, font="Arial")
            if s.get("setup"):
                for chunk in re.split(r"\n\s*\n", str(s["setup"]).strip()):
                    runs(c0.add_paragraph(), " ".join(chunk.split()), size=9.5)
            if s.get("listen_for"):
                pl = c0.add_paragraph(); runs(pl, "LISTEN FOR", size=7, bold=True, color=MUTED, font="Arial")
                for x in s["listen_for"]:
                    runs(c0.add_paragraph(), "•  " + x, size=9.5)
            runs(c0.add_paragraph(), s["url"], size=8, color=MUTED, font="Arial")
            pc = c1.paragraphs[0]
            pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pc.add_run().add_picture(qr_png(s["url"]), width=Inches(1.05))
            cap = c1.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs(cap, "Scan to watch", size=7, color=MUTED, font="Arial")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "questions":
            if s.get("body"): paras(s["body"])
            numbered(s.get("items") or [], s.get("note_lines", 0))
        elif t == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            border(p, "left", "93371F", 12)
            runs(p, " ".join(str(s["text"]).split()), italic=True, size=11)
            if s.get("source"):
                runs(doc.add_paragraph(), "— " + s["source"], size=8.5, color=MUTED, font="Arial").paragraph_format.left_indent = Inches(0.2)
        elif t == "activity":
            if s.get("body"): paras(s["body"])
            if s.get("steps"): numbered(s["steps"])

    if L.get("handout_notes_lines"):
        heading("Notes")
        lines(L["handout_notes_lines"])

    # footer block
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    border(p, "top")
    tbl = doc.add_table(rows=1, cols=2)
    widths(tbl, [1.0, 5.9])
    a, b = tbl.rows[0].cells
    a.paragraphs[0].add_run().add_picture(qr_png(site["url"]), width=Inches(0.8))
    pb = b.paragraphs[0]
    runs(pb, "Watch the clips and see every lesson online: ", size=8.5, bold=True, font="Arial")
    runs(pb, site["url"].replace("https://", ""), size=8.5, font="Arial", color=ACC)
    runs(b.add_paragraph(), f"{site['class_name']} · {site['teacher']} · " + " · ".join(site["meetings"]), size=8, color=MUTED, font="Arial")
    runs(b.add_paragraph(), f"{site['church']['name']}, {site['church']['address']}", size=8, color=MUTED, font="Arial")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))

# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--preview", action="store_true", help="include drafts; write to preview/ (not deployed)")
    ap.add_argument("--no-handouts", action="store_true")
    a = ap.parse_args()
    out = ROOT / ("preview" if a.preview else "public")
    out.mkdir(exist_ok=True)
    site, lessons = load(a.preview)
    print(f"Building {'PREVIEW' if a.preview else 'LIVE'} site -> {out.relative_to(ROOT)}/")
    for L in lessons:
        print(f"  week {L['week']}: {L['status']:<9} {L.get('title','')}")
    build_site(site, lessons, out, a.preview)
    if not a.no_handouts:
        build_handouts(site, lessons, out)
    print("Done.")

if __name__ == "__main__":
    main()
